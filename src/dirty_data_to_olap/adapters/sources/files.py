"""Read-only CSV, Parquet and optional XLSX source adapters.

Parsing and columnar operations remain inside this concrete adapter family.
The public methods return only project-owned contracts and staged references.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence

import pyarrow as pa
import pyarrow.parquet as pq

from dirty_data_to_olap.adapters.sources.staging import SourceFaithfulParquetStager
from dirty_data_to_olap.application.source_adapter import SourceAdapter
from dirty_data_to_olap.domain.contracts.source import (
    AdapterReference,
    ColumnDescriptor,
    DeclaredConstraint,
    ExtractionMetrics,
    MaxRowsScope,
    ObservationMode,
    ObservationScope,
    RecordLocatorKind,
    RowAccounting,
    RowCountSemantics,
    SelectionScope,
    SnapshotConsistency,
    SourceCatalog,
    SourceDescriptor,
    SourceFailure,
    SourceFailureKind,
    SourceIngestionError,
    SourceRecordReference,
    SourceRegistryRecord,
    SourceSelection,
    SourceSnapshot,
    SourceSnapshotResult,
    SourceTableKind,
    SourceType,
    StabilityScope,
    TableDescriptor,
    TableObservationStatus,
    TableSnapshotObservation,
    column_id_for,
    file_content_fingerprint,
    normalized_file_locator,
    record_ref_for_ordinal,
    record_ref_for_pk,
    schema_fingerprint,
    snapshot_id_for,
    source_id_for,
    stable_digest,
    table_id_for,
    utc_now,
)


def _failure(kind: SourceFailureKind, operation: str, detail: str, *, observed: int = 0) -> SourceIngestionError:
    return SourceIngestionError(SourceFailure(
        kind=kind,
        operation=operation,
        detail=detail,
        retryable=kind in {SourceFailureKind.ACCESS_FAILED, SourceFailureKind.TIMEOUT, SourceFailureKind.STAGING_FAILED},
        records_observed_before_failure=observed,
    ))


def detect_file_source_type(locator: str | Path) -> SourceType:
    suffix = Path(locator).suffix.lower()
    try:
        return {".csv": SourceType.CSV, ".parquet": SourceType.PARQUET, ".xlsx": SourceType.XLSX}[suffix]
    except KeyError:
        raise _failure(SourceFailureKind.UNSUPPORTED, "detect_file_source_type", "file extension is not a supported V1 tabular source") from None


def _effective_scope(record: SourceRegistryRecord, selection: SourceSelection) -> SelectionScope:
    selected = selection.scope
    configured = record.scope
    return SelectionScope(
        included_objects=selected.included_objects or configured.included_objects,
        excluded_objects=selected.excluded_objects or configured.excluded_objects,
        include_views=selected.include_views or configured.include_views,
        included_columns=selected.included_columns or configured.included_columns,
    )


def _adapter_reference(name: str, version: str, record: SourceRegistryRecord, scope: SelectionScope) -> AdapterReference:
    return AdapterReference(
        name=name,
        version=version,
        config_fingerprint=stable_digest({
            "adapter_config": dict(record.adapter_config),
            "scope": scope,
        }),
    )


def _schema_type_from_arrow(data_type: pa.DataType) -> str:
    if pa.types.is_boolean(data_type):
        return "boolean"
    if pa.types.is_integer(data_type):
        return "integer"
    if pa.types.is_floating(data_type):
        return "real"
    if pa.types.is_decimal(data_type):
        return "decimal"
    if pa.types.is_date(data_type):
        return "date"
    if pa.types.is_timestamp(data_type):
        return "timestamp"
    if pa.types.is_binary(data_type):
        return "binary"
    if pa.types.is_list(data_type) or pa.types.is_struct(data_type) or pa.types.is_map(data_type):
        return "nested"
    return "text"


def _table_catalog(
    *,
    source_id: str,
    source_type: SourceType,
    physical_name: str,
    kind: SourceTableKind,
    columns: Sequence[tuple[str, int, str, str, bool | None, bool | None, str | None]],
    row_count: int | None,
    row_count_semantics: RowCountSemantics,
    adapter_reference: AdapterReference,
) -> tuple[TableDescriptor, tuple[ColumnDescriptor, ...], tuple[DeclaredConstraint, ...]]:
    table_id = table_id_for(source_id, None, physical_name, kind)
    table = TableDescriptor(
        table_id=table_id,
        source_id=source_id,
        physical_name=physical_name,
        table_kind=kind,
        row_count=row_count,
        row_count_semantics=row_count_semantics,
        fingerprint=stable_digest({"name": physical_name, "kind": kind.value, "columns": columns}),
    )
    descriptors = tuple(
        ColumnDescriptor(
            column_id=column_id_for(table_id, name, ordinal),
            table_id=table_id,
            physical_name=name,
            ordinal=ordinal,
            native_physical_type=native_type,
            normalized_physical_type=normalized_type,
            schema_nullable=nullable,
            raw_declared_notnull=raw_notnull,
            declared_default=default,
        )
        for name, ordinal, native_type, normalized_type, nullable, raw_notnull, default in columns
    )
    return table, descriptors, ()


def _make_catalog(
    *,
    record: SourceRegistryRecord,
    selection: SourceSelection,
    source_id: str,
    source_type: SourceType,
    source_locator: str,
    physical_name: str,
    kind: SourceTableKind,
    columns: Sequence[tuple[str, int, str, str, bool | None, bool | None, str | None]],
    row_count: int | None,
    row_count_semantics: RowCountSemantics,
    source_fingerprint: str,
    adapter_reference: AdapterReference,
) -> SourceCatalog:
    table, column_descriptors, constraints = _table_catalog(
        source_id=source_id,
        source_type=source_type,
        physical_name=physical_name,
        kind=kind,
        columns=columns,
        row_count=row_count,
        row_count_semantics=row_count_semantics,
        adapter_reference=adapter_reference,
    )
    schema_fp = schema_fingerprint({
        "source_type": source_type.value,
        "selection": selection.scope,
        "tables": [table],
        "columns": column_descriptors,
        "constraints": constraints,
    })
    descriptor = SourceDescriptor(
        source_id=source_id,
        display_name=record.display_name,
        source_type=source_type,
        file_locator=source_locator,
        selection_scope=selection.scope,
        source_fingerprint=source_fingerprint,
        schema_fingerprint=schema_fp,
        adapter_reference=adapter_reference,
    )
    return SourceCatalog(
        source=descriptor,
        tables=(table,),
        columns=column_descriptors,
        declared_constraints=constraints,
    )


def _record_references(
    *,
    source_id: str,
    snapshot_id: str,
    table_id: str,
    batch_id: str,
    rows: Sequence[Mapping[str, Any]],
    first_ordinal: int,
    columns: Sequence[ColumnDescriptor],
) -> tuple[SourceRecordReference, ...]:
    key_columns = tuple(
        column.physical_name
        for column in sorted(
            (item for item in columns if item.primary_key_position is not None),
            key=lambda item: item.primary_key_position or 0,
        )
    )
    refs: list[SourceRecordReference] = []
    for offset, row in enumerate(rows):
        ordinal = first_ordinal + offset
        if key_columns:
            key_values = tuple(row.get(key) for key in key_columns)
            ref = SourceRecordReference(
                record_ref=record_ref_for_pk(source_id, snapshot_id, table_id, key_columns, key_values),
                source_id=source_id,
                snapshot_id=snapshot_id,
                table_id=table_id,
                batch_id=batch_id,
                extraction_ordinal=ordinal,
                locator_kind=RecordLocatorKind.PRIMARY_KEY,
                key_columns=key_columns,
                key_values=key_values,
                stability_scope=StabilityScope.SOURCE_KEY_SCOPE,
            )
        else:
            ref = SourceRecordReference(
                record_ref=record_ref_for_ordinal(source_id, snapshot_id, table_id, ordinal),
                source_id=source_id,
                snapshot_id=snapshot_id,
                table_id=table_id,
                batch_id=batch_id,
                extraction_ordinal=ordinal,
                locator_kind=RecordLocatorKind.SNAPSHOT_ORDINAL,
                stability_scope=StabilityScope.SNAPSHOT_ONLY,
            )
        refs.append(ref)
    return tuple(refs)


class FileSourceAdapter(SourceAdapter):
    """One coherent adapter family for CSV, Parquet and optional XLSX."""

    name = "file_source"
    version = "1.0.0"

    def __init__(self, source_type: SourceType, *, project_root: Path) -> None:
        if source_type not in {SourceType.CSV, SourceType.PARQUET, SourceType.XLSX}:
            raise ValueError("FileSourceAdapter requires csv, parquet or xlsx")
        self.source_type = source_type
        self.project_root = project_root.resolve()

    def _confined_path(self, locator: str | Path, *, operation: str) -> Path:
        try:
            path = Path(locator).expanduser().resolve(strict=True)
            path.relative_to(self.project_root)
        except (OSError, ValueError) as exc:
            raise _failure(SourceFailureKind.ACCESS_FAILED, operation, "source file is outside the project-owned source root") from exc
        if not path.is_file():
            raise _failure(SourceFailureKind.ACCESS_FAILED, operation, "source file is not accessible")
        return path

    def _path(self, record: SourceRegistryRecord) -> Path:
        if not record.file_locator:
            raise _failure(SourceFailureKind.INVALID_SELECTION, "resolve_file", "file locator is missing")
        return self._confined_path(record.file_locator, operation="resolve_file")

    def _object_name(self, path: Path, scope: SelectionScope) -> str:
        if self.source_type is SourceType.XLSX:
            return (scope.included_objects[0] if scope.included_objects else "Sheet1")
        return path.stem

    def discover_source(self, selection: SourceSelection, registry_record: SourceRegistryRecord) -> SourceCatalog:
        path = self._path(registry_record)
        if detect_file_source_type(path) is not self.source_type:
            raise _failure(SourceFailureKind.UNSUPPORTED, "discover_file_source", "file extension does not match the selected source adapter")
        locator = normalized_file_locator(path)
        source_id = registry_record.source_id or source_id_for(self.source_type, locator)
        scope = _effective_scope(registry_record, selection)
        reference = _adapter_reference(self.name, self.version, registry_record, scope)
        if self.source_type is SourceType.CSV:
            columns = self._discover_csv(path)
            name = self._object_name(path, scope)
            count = None
            semantics = RowCountSemantics.UNKNOWN
        elif self.source_type is SourceType.PARQUET:
            name, columns, count, semantics = self._discover_parquet(path, scope)
        else:
            name, columns, count, semantics = self._discover_xlsx(path, scope)
        catalog_selection = selection.model_copy(update={"scope": scope})
        return _make_catalog(
            record=registry_record,
            selection=catalog_selection,
            source_id=source_id,
            source_type=self.source_type,
            source_locator=locator,
            physical_name=name,
            kind=SourceTableKind.WORKSHEET if self.source_type is SourceType.XLSX else SourceTableKind.FILE,
            columns=columns,
            row_count=count,
            row_count_semantics=semantics,
            source_fingerprint=file_content_fingerprint(path),
            adapter_reference=reference,
        )

    def _discover_csv(self, path: Path) -> tuple[tuple[str, int, str, str, bool | None, bool | None, str | None], ...]:
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, strict=True)
                header = next(reader, None)
        except (OSError, UnicodeError, csv.Error):
            raise _failure(SourceFailureKind.PARSE_FAILED, "discover_csv", "CSV header could not be parsed") from None
        if not header or any(not name for name in header) or len(set(header)) != len(header):
            raise _failure(SourceFailureKind.PARSE_FAILED, "discover_csv", "CSV header must contain unique non-empty names")
        return tuple((name, ordinal, "TEXT", "text", None, None, None) for ordinal, name in enumerate(header))

    def _discover_parquet(self, path: Path, scope: SelectionScope) -> tuple[str, tuple[tuple[str, int, str, str, bool | None, bool | None, str | None], ...], int, RowCountSemantics]:
        try:
            parquet_file = pq.ParquetFile(path)
            schema = parquet_file.schema_arrow
            name = self._object_name(path, scope)
            columns = tuple((field.name, ordinal, str(field.type), _schema_type_from_arrow(field.type), field.nullable, None, None) for ordinal, field in enumerate(schema))
            return name, columns, parquet_file.metadata.num_rows, RowCountSemantics.EXACT
        except Exception:
            raise _failure(SourceFailureKind.PARSE_FAILED, "discover_parquet", "Parquet schema could not be inspected") from None

    def _discover_xlsx(self, path: Path, scope: SelectionScope) -> tuple[str, tuple[tuple[str, int, str, str, bool | None, bool | None, str | None], ...], None, RowCountSemantics]:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                sheet_name = self._object_name(path, scope)
                if sheet_name not in workbook.sheetnames:
                    raise ValueError("selected worksheet was not found")
                sheet = workbook[sheet_name]
                header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
                names = tuple(str(value) if value is not None else "" for value in (header or ()))
            finally:
                workbook.close()
        except Exception:
            raise _failure(SourceFailureKind.PARSE_FAILED, "discover_xlsx", "XLSX worksheet header could not be parsed") from None
        if not names or any(not name for name in names) or len(set(names)) != len(names):
            raise _failure(SourceFailureKind.PARSE_FAILED, "discover_xlsx", "XLSX header must contain unique non-empty names")
        columns = tuple((name, ordinal, "CELL", "unknown", True, None, None) for ordinal, name in enumerate(names))
        return sheet_name, columns, None, RowCountSemantics.UNKNOWN

    def _rows(self, path: Path, table: TableDescriptor, columns: Sequence[ColumnDescriptor], selection: SourceSelection) -> tuple[Iterator[tuple[int, Mapping[str, Any]]], pa.Schema | None]:
        chosen = selection.scope.included_columns.get(table.physical_name)
        names = tuple(column.physical_name for column in columns if not chosen or column.physical_name in chosen)
        if self.source_type is SourceType.CSV:
            def csv_rows() -> Iterator[tuple[int, Mapping[str, Any]]]:
                try:
                    with path.open("r", encoding="utf-8-sig", newline="") as stream:
                        reader = csv.reader(stream, strict=True)
                        header = next(reader, None)
                        if header is None:
                            raise _failure(SourceFailureKind.PARSE_FAILED, "extract_csv", "CSV has no header")
                        positions = [header.index(name) for name in names]
                        for ordinal, values in enumerate(reader):
                            if len(values) != len(header):
                                raise _failure(SourceFailureKind.PARSE_FAILED, "extract_csv", "malformed CSV row was rejected explicitly", observed=ordinal)
                            yield ordinal, {name: values[position] for name, position in zip(names, positions)}
                except SourceIngestionError:
                    raise
                except (OSError, UnicodeError, csv.Error):
                    raise _failure(SourceFailureKind.PARSE_FAILED, "extract_csv", "CSV row parsing failed") from None
            return csv_rows(), pa.schema([pa.field(name, pa.string(), nullable=True) for name in names])
        if self.source_type is SourceType.PARQUET:
            def parquet_rows() -> Iterator[tuple[int, Mapping[str, Any]]]:
                try:
                    parquet_file = pq.ParquetFile(path)
                    ordinal = 0
                    for batch in parquet_file.iter_batches(batch_size=selection.extraction.chunk_size, columns=list(names)):
                        for row in batch.to_pylist():
                            yield ordinal, row
                            ordinal += 1
                except Exception:
                    raise _failure(SourceFailureKind.EXTRACTION_FAILED, "extract_parquet", "Parquet batch iteration failed") from None
            schema = pq.ParquetFile(path).schema_arrow
            return parquet_rows(), schema if names == tuple(schema.names) else pa.schema([schema.field(name) for name in names])
        def xlsx_rows() -> Iterator[tuple[int, Mapping[str, Any]]]:
            try:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    sheet = workbook[table.physical_name]
                    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
                    all_names = tuple(str(value) if value is not None else "" for value in (header or ()))
                    positions = [all_names.index(name) for name in names]
                    for ordinal, values in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
                        if not any(value is not None for value in values):
                            continue
                        padded = tuple(values) + (None,) * max(0, len(all_names) - len(values))
                        yield ordinal, {name: padded[position] for name, position in zip(names, positions)}
                finally:
                    workbook.close()
            except SourceIngestionError:
                raise
            except Exception:
                raise _failure(SourceFailureKind.EXTRACTION_FAILED, "extract_xlsx", "XLSX row iteration failed") from None
        return xlsx_rows(), None

    def create_bounded_snapshot(
        self,
        catalog: SourceCatalog,
        selection: SourceSelection,
        *,
        execution_context_id: str,
        staging_root: Path,
    ) -> SourceSnapshotResult:
        path = self._confined_path(catalog.source.file_locator or "", operation="resolve_snapshot_file")
        table = catalog.tables[0]
        columns = tuple(column for column in catalog.columns if column.table_id == table.table_id)
        snapshot_id = snapshot_id_for(catalog.source_id, catalog.source.schema_fingerprint, catalog.source.selection_scope, selection.extraction, execution_context_id)
        snapshot = SourceSnapshot(
            source_id=catalog.source_id,
            snapshot_id=snapshot_id,
            execution_context_id=execution_context_id,
            schema_fingerprint=catalog.source.schema_fingerprint,
            source_fingerprint=catalog.source.source_fingerprint,
            observed_at=utc_now(),
            selection_scope=catalog.source.selection_scope,
            observation_scope=ObservationScope(
                mode=ObservationMode.BOUNDED if selection.extraction.max_rows is not None else ObservationMode.FULL,
                chunk_size=selection.extraction.chunk_size,
                max_rows=selection.extraction.max_rows,
                max_rows_scope=selection.extraction.max_rows_scope,
                input_records_observed=0,
            ),
            extraction_policy=selection.extraction,
            consistency=SnapshotConsistency.FILE_IMMUTABLE,
            adapter_reference=catalog.source.adapter_reference,
        )
        stager = SourceFaithfulParquetStager(self.project_root, adapter_reference=catalog.source.adapter_reference)
        extraction_selection = selection.model_copy(update={"scope": catalog.source.selection_scope})
        rows_iter, schema = self._rows(path, table, columns, extraction_selection)
        batches = []
        references = []
        table_observations = []
        observed = 0
        staged = 0
        pending: list[Mapping[str, Any]] = []
        pending_first = 0
        batch_index = 0
        try:
            expected_fingerprint = catalog.source.source_fingerprint
            if expected_fingerprint is None:
                raise _failure(SourceFailureKind.SNAPSHOT_INVALID, "verify_file_before_extraction", "file snapshot has no discovery fingerprint")
            before_fingerprint = file_content_fingerprint(path)
            if before_fingerprint != expected_fingerprint:
                raise _failure(SourceFailureKind.SNAPSHOT_INVALID, "verify_file_before_extraction", "file content changed after discovery")
            table_rows_observed = 0
            table_exhausted = True
            for ordinal, row in rows_iter:
                if selection.extraction.max_rows is not None and observed >= selection.extraction.max_rows:
                    table_exhausted = False
                    break
                if not pending:
                    pending_first = ordinal
                pending.append(row)
                observed += 1
                table_rows_observed += 1
                if len(pending) >= selection.extraction.chunk_size:
                    batch = stager.stage_rows(
                        source_id=catalog.source_id, snapshot_id=snapshot_id, table_id=table.table_id,
                        batch_index=batch_index, first_ordinal=pending_first, rows=pending,
                        schema_fingerprint=catalog.source.schema_fingerprint, schema=schema, staging_root=staging_root,
                    )
                    batches.append(batch)
                    references.extend(_record_references(source_id=catalog.source_id, snapshot_id=snapshot_id, table_id=table.table_id, batch_id=batch.batch_id, rows=pending, first_ordinal=pending_first, columns=columns))
                    staged += len(pending)
                    pending = []
                    batch_index += 1
            if pending:
                batch = stager.stage_rows(
                    source_id=catalog.source_id, snapshot_id=snapshot_id, table_id=table.table_id,
                    batch_index=batch_index, first_ordinal=pending_first, rows=pending,
                    schema_fingerprint=catalog.source.schema_fingerprint, schema=schema, staging_root=staging_root,
                )
                batches.append(batch)
                references.extend(_record_references(source_id=catalog.source_id, snapshot_id=snapshot_id, table_id=table.table_id, batch_id=batch.batch_id, rows=pending, first_ordinal=pending_first, columns=columns))
                staged += len(pending)
            if table_exhausted:
                table_status = TableObservationStatus.FULLY_OBSERVED
            elif table_rows_observed:
                table_status = TableObservationStatus.PARTIALLY_OBSERVED
            else:
                table_status = TableObservationStatus.NOT_OBSERVED
            table_observations.append(TableSnapshotObservation(table_id=table.table_id, rows_observed=table_rows_observed, status=table_status))
            after_fingerprint = file_content_fingerprint(path)
            if after_fingerprint != before_fingerprint:
                raise _failure(SourceFailureKind.SNAPSHOT_INVALID, "verify_file_after_extraction", "file content changed during extraction")
        except SourceIngestionError:
            stager.discard_batches(batches)
            raise
        accounting = RowAccounting(
            input_records_observed=observed,
            successfully_staged_records=staged,
            explicitly_quarantined_records=0,
            unresolved_records=0,
            accounting_complete=True,
        )
        result = SourceSnapshotResult(
            snapshot=snapshot.model_copy(update={"observation_scope": snapshot.observation_scope.model_copy(update={"input_records_observed": observed})}),
            batches=tuple(batches),
            record_references=tuple(references),
            accounting=accounting,
            metrics=ExtractionMetrics(input_records_observed=observed, staged_records=staged, batch_count=len(batches), configured_chunk_size=selection.extraction.chunk_size, bytes_staged=sum(Path(self.project_root / batch.artifact_location).stat().st_size for batch in batches)),
            table_observations=tuple(table_observations),
        )
        stager.write_catalog(catalog, run_root=staging_root.parent)
        stager.write_manifest(result, run_root=staging_root.parent)
        return result
